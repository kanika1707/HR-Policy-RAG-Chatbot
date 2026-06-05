"""
streamlit_app.py — Chat UI for the HR Policy RAG Chatbot.

Streamlit re-runs this entire script from top to bottom every time the user
sends a message or clicks a button.  Persistent state (chat history, index
ready-flag, etc.) is kept in st.session_state, which survives re-runs.

Layout:
  Left sidebar  → document list, Rebuild Index, Export Chat, pipeline info
  Main area     → chat history + chat input box at the bottom
"""

import io                              # used to create an in-memory file for DOCX export
import streamlit as st
from pathlib import Path
from docx import Document as DocxDocument  # python-docx — creates Word files

from rag_graph import search, build_index, init_index  # our RAG pipeline

# ── Constants ─────────────────────────────────────────────────────────────────
DATA_DIR     = Path(__file__).parent / "data"   # folder where documents live
SUPPORTED    = {".txt", ".pdf", ".docx"}        # accepted file extensions
MEMORY_TURNS = 6   # how many past messages to include as conversational context


# ─────────────────────────────────────────────────────────────────────────────
# Helper: format conversation history for the LLM prompt (Step H)
# ─────────────────────────────────────────────────────────────────────────────
def _format_history(messages: list) -> str:
    """
    Take the last MEMORY_TURNS messages from the chat history and format
    them as a readable block that gets prepended to the LLM prompt.

    Example output:
        Chat History:
        User: what is the leave policy?
        Assistant: The leave policy covers ...
        User: how many days for sick leave?
        Assistant: Employees get 10 days ...

    This lets the LLM understand follow-up questions like "can you elaborate?"
    """
    recent = messages[-MEMORY_TURNS:]
    if not recent:
        return ""   # no history on first message

    lines = ["Chat History:"]
    for m in recent:
        role = "User" if m["role"] == "user" else "Assistant"
        lines.append(f"{role}: {m['content']}")

    # Trailing newlines separate history block from the Context block in the prompt
    return "\n".join(lines) + "\n\n"


# ─────────────────────────────────────────────────────────────────────────────
# Helper: export the full conversation to a Word document (Step I)
# ─────────────────────────────────────────────────────────────────────────────
def _export_docx(messages: list) -> io.BytesIO:
    """
    Convert the chat history list into a .docx file held in memory.
    Returns an io.BytesIO buffer that Streamlit's download_button can serve.

    Each message is written as:
        You:        <user text>
        Assistant:  <bot text>
    """
    doc = DocxDocument()
    doc.add_heading("Chat Export", level=0)   # big title at the top

    for m in messages:
        role = "You" if m["role"] == "user" else "Assistant"
        p = doc.add_paragraph()
        p.add_run(f"{role}: ").bold = True    # bold speaker label
        p.add_run(m["content"])               # message text (normal weight)

    # Save to an in-memory buffer instead of writing a real file
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)   # rewind so Streamlit can read from the start
    return buf


# ─────────────────────────────────────────────────────────────────────────────
# Page config — must be the very first Streamlit call
# ─────────────────────────────────────────────────────────────────────────────
st.set_page_config(page_title="RAG Chatbot", page_icon="🤖", layout="centered")
st.title("🤖 RAG Chatbot")
st.caption("Ask questions about your documents. Powered by Groq · Llama-3.3-70b")


# ─────────────────────────────────────────────────────────────────────────────
# Sidebar
# ─────────────────────────────────────────────────────────────────────────────
with st.sidebar:

    # Show which documents are currently indexed
    st.header("📁 Documents")
    files = [f for f in DATA_DIR.rglob("*") if f.suffix.lower() in SUPPORTED]
    if files:
        st.success(f"{len(files)} document(s) indexed")
        for f in files:
            st.text(f"• {f.name}")
    else:
        st.warning("No documents found.\nAdd .txt / .pdf / .docx files to data/")

    st.divider()

    # Rebuild Index button — use this after adding or removing documents
    # force=True tells build_index to delete and recreate the FAISS index
    if st.button("🔄 Rebuild Index", use_container_width=True):
        with st.spinner("Rebuilding index..."):
            try:
                build_index(force=True)    # re-chunk and re-embed all documents
                init_index(force=True)     # reload the chain with the new index
                st.success("Index rebuilt!")
                st.rerun()                 # refresh the page to show updated state
            except Exception as e:
                st.error(str(e))

    st.divider()

    # Export Chat button — only shown when there are messages to export
    if st.session_state.get("messages"):
        st.download_button(
            label="📄 Export Chat (.docx)",
            data=_export_docx(st.session_state.messages),
            file_name="chat_export.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )

    # Clear Chat button — wipes session_state.messages and reloads the page
    if st.button("🗑️ Clear Chat", use_container_width=True):
        st.session_state.messages = []
        st.rerun()

    st.divider()

    # Pipeline summary shown at the bottom of the sidebar for quick reference
    st.markdown(
        "**Pipeline**\n"
        "- Semantic + Hierarchical chunking (800 chars)\n"
        "- all-MiniLM-L6-v2 embeddings\n"
        "- FAISS dense retrieval (k=10)\n"
        "- BM25 reranking (top 5)\n"
        "- Groq Llama-3.3-70b\n"
        "- Conversational memory\n"
        "- DOCX export"
    )


# ─────────────────────────────────────────────────────────────────────────────
# Lazy index initialisation
# ─────────────────────────────────────────────────────────────────────────────
# "ready" flag is stored in session_state so we only do this once per session.
# On first load this will download the embedding model and build the FAISS index
# (~30 seconds).  On subsequent page re-runs (caused by user input) it is a
# no-op because "ready" is already True.
if "ready" not in st.session_state:
    with st.spinner("Loading index... (first run builds embeddings, ~30 s)"):
        try:
            init_index()                        # build or load from disk
            st.session_state.ready = True       # mark as done for future re-runs
        except Exception as e:
            st.error(
                f"Failed to load index: {e}\n\n"
                "Add documents to data/ then click Rebuild Index."
            )
            st.stop()   # halt the script — nothing else should render


# ─────────────────────────────────────────────────────────────────────────────
# Helper: render referenced chunks in an expander below an answer
# ─────────────────────────────────────────────────────────────────────────────
def _render_chunks(chunks: list):
    """
    Display each retrieved chunk in a collapsible expander.

    For each chunk shows:
      - Chunk number
      - Filename as a clickable link that opens the file locally
      - Page number (PDFs only)
      - The exact text the LLM used, styled in a blue-left-bordered box
    """
    if not chunks:
        return

    with st.expander(f"📄 Referenced chunks ({len(chunks)})"):
        for i, chunk in enumerate(chunks, 1):
            fname = chunk["filename"] or "Unknown"
            page  = chunk.get("page")
            uri   = chunk.get("file_uri", "")

            # Build the header line: "Chunk 1 — [filename.pdf](file:///...)"
            header = f"**Chunk {i} —** "
            if uri:
                # Markdown link using file:// URI → clicks open the file in browser
                header += f"[{fname}]({uri})"
            else:
                header += f"`{fname}`"

            # Add page number for PDFs (LangChain uses 0-based page index, display as 1-based)
            if page is not None:
                header += f"  •  page {int(page) + 1}"

            st.markdown(header, unsafe_allow_html=True)

            # Render the chunk text in a styled blockquote-style div
            st.markdown(
                f"<div style='background:#f0f2f6;border-left:3px solid #4a90d9;"
                f"padding:8px 12px;border-radius:4px;font-size:0.85em;"
                f"white-space:pre-wrap'>{chunk['content']}</div>",
                unsafe_allow_html=True,
            )

            # Add a divider between chunks (but not after the last one)
            if i < len(chunks):
                st.divider()


# ─────────────────────────────────────────────────────────────────────────────
# Chat history display
# ─────────────────────────────────────────────────────────────────────────────
# Initialise the messages list on the very first run
if "messages" not in st.session_state:
    st.session_state.messages = []

# Replay all previous messages so the user sees the full conversation
for msg in st.session_state.messages:
    with st.chat_message(msg["role"]):           # "user" → human icon, "assistant" → bot icon
        st.markdown(msg["content"])              # render markdown in answers
        if msg.get("chunks"):
            _render_chunks(msg["chunks"])        # show source expander
        if msg.get("meta"):
            st.caption(msg["meta"])              # show query-type caption in small text


# ─────────────────────────────────────────────────────────────────────────────
# Chat input — this is the main interaction loop
# ─────────────────────────────────────────────────────────────────────────────
# st.chat_input returns None when empty, or the question string when submitted.
# The walrus operator (:=) assigns and checks in one step.
if question := st.chat_input("Ask anything about your documents..."):

    # 1. Save user message and display it immediately
    st.session_state.messages.append({"role": "user", "content": question})
    with st.chat_message("user"):
        st.markdown(question)

    # 2. Generate the answer using the RAG pipeline
    with st.chat_message("assistant"):
        with st.spinner("Thinking..."):
            try:
                # Build the conversation history string from all messages
                # except the one just appended (we pass history, not the current question)
                history = _format_history(st.session_state.messages[:-1])

                # Call the full RAG pipeline: retrieve → rerank → LLM
                result = search(question, chat_history=history)
                answer = result["answer"]
                chunks = result["chunks"]
                qtype  = result["query_type"]

                # Caption text shown below the answer in small grey text
                meta = (
                    f"Query: **{qtype}** | "
                    f"{'parent chunks' if qtype == 'broad' else 'child chunks'} + BM25 rerank"
                )
            except Exception as e:
                answer, chunks, meta = f"Error: {e}", [], ""

        # 3. Display the answer, source chunks, and metadata caption
        st.markdown(answer)
        _render_chunks(chunks)
        if meta:
            st.caption(meta)

    # 4. Persist the assistant message for future re-runs
    st.session_state.messages.append({
        "role":    "assistant",
        "content": answer,
        "chunks":  chunks,   # saved so source expanders work when scrolling up
        "meta":    meta,
    })
